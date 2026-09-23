from __future__ import annotations

from io import BytesIO

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from app.models import Meeting


def _set_run_font(run, name: str = "Arial", size: int = 11, bold: bool = False) -> None:
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)


def _shade_cell(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def _set_cell_margins(cell, margin: int = 110) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for edge in ("top", "start", "bottom", "end"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:w"), str(margin))
        element.set(qn("w:type"), "dxa")
        margins.append(element)


def build_protocol_docx(meeting: Meeting) -> BytesIO:
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(10)
    title.paragraph_format.space_after = Pt(18)
    _set_run_font(title.add_run("Протокол совещания"), size=18, bold=True)

    metadata = document.add_paragraph()
    metadata.paragraph_format.space_after = Pt(4)
    _set_run_font(metadata.add_run("Тема: "), bold=True)
    _set_run_font(metadata.add_run(meeting.title))
    metadata = document.add_paragraph()
    metadata.paragraph_format.space_after = Pt(14)
    _set_run_font(metadata.add_run("Дата: "), bold=True)
    _set_run_font(metadata.add_run(meeting.date.strftime("%d.%m.%Y %H:%M")))

    participant_by_label = {item.speaker_label: item for item in meeting.participants}
    heading = document.add_heading("Текст совещания", level=1)
    _set_run_font(heading.runs[0], size=14, bold=True)
    for segment in meeting.segments:
        participant = participant_by_label.get(segment.speaker_label)
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(8)
        speaker = participant.name if participant else segment.speaker_label or "Говорящий не определён"
        role = f", {participant.role}" if participant else ""
        _set_run_font(paragraph.add_run(f"{speaker}{role} [{int(segment.start // 60):02d}:{int(segment.start % 60):02d}]\n"), bold=True)
        _set_run_font(paragraph.add_run(segment.text))

    heading = document.add_heading("Саммари", level=1)
    _set_run_font(heading.runs[0], size=14, bold=True)
    summary = document.add_paragraph()
    summary.paragraph_format.space_after = Pt(14)
    _set_run_font(summary.add_run(meeting.summary or "Саммари ещё не сформировано."))

    heading = document.add_heading("Поручения", level=1)
    _set_run_font(heading.runs[0], size=14, bold=True)
    table = document.add_table(rows=1, cols=3)
    table.autofit = False
    table.columns[0].width = Inches(3.7)
    table.columns[1].width = Inches(1.65)
    table.columns[2].width = Inches(1.25)
    table.style = "Table Grid"
    headers = ("Поручение", "Ответственный", "Срок")
    for index, label in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        _shade_cell(cell, "17365D")
        _set_cell_margins(cell)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = cell.paragraphs[0].add_run(label)
        _set_run_font(run, size=10, bold=True)
        run.font.color.rgb = RGBColor(255, 255, 255)
    for row_index, action in enumerate(meeting.action_items):
        cells = table.add_row().cells
        values = (
            action.task,
            action.assignee or "Не определён",
            action.deadline_date.strftime("%d.%m.%Y") if action.deadline_date else action.deadline_raw or "Не указан",
        )
        for index, value in enumerate(values):
            cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _set_cell_margins(cells[index])
            if row_index % 2:
                _shade_cell(cells[index], "F3F6FA")
            paragraph = cells[index].paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if index < 2 else WD_ALIGN_PARAGRAPH.CENTER
            _set_run_font(paragraph.add_run(value), size=10)

    output = BytesIO()
    document.save(output)
    output.seek(0)
    return output
